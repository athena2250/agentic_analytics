"""
The event-analysis workbook (plan §17.5).

"Ready to use" is treated here as a testable property rather than a stylistic
one: a frozen bold header row, columns wide enough to read, real Excel number
formats on amounts and percentages, no index column, and sheet names Excel will
actually accept. Deltas are written as numbers, so a reader can sort by them;
percentages are written as fractions carrying a percent format, so 0.12 renders
as 12.0% and still sorts as 0.12.

The Method sheet is the reason the rest can be forwarded. A workbook that
leaves the conversation has to disclose what the system assumed: which column
filled which role and whether it was inferred or confirmed, the event and
baseline windows, every query that ran, and every sub-analysis that did not.
"""
from __future__ import annotations

import io
import re

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

# Excel's own limits, not preferences: a sheet name may not exceed 31
# characters and may not contain any of these.
_MAX_SHEET_NAME = 31
_ILLEGAL_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")

# Column widths are clamped so one long free-text cell can't push a column off
# the screen, and so a narrow column still shows its header.
_MIN_WIDTH, _MAX_WIDTH = 10, 60

_NUMBER_FORMATS = {
    "count": "#,##0",
    "amount": "#,##0.00",
    "ratio": "#,##0.00",
    # Stored as a fraction; Excel renders the percent sign. Pre-multiplying by
    # 100 and appending "%" as text would make the column unsortable.
    "pct": "0.0%",
}


def safe_sheet_name(name: str, used: set[str]) -> str:
    """An Excel-legal, unique sheet name. Truncation keeps the front of the
    name, and a collision gets a numeric suffix rather than overwriting."""
    cleaned = _ILLEGAL_SHEET_CHARS.sub("-", str(name)).strip() or "Sheet"
    cleaned = cleaned[:_MAX_SHEET_NAME]
    if cleaned not in used:
        used.add(cleaned)
        return cleaned
    for n in range(2, 100):
        suffix = f" {n}"
        candidate = cleaned[: _MAX_SHEET_NAME - len(suffix)] + suffix
        if candidate not in used:
            used.add(candidate)
            return candidate
    used.add(cleaned)
    return cleaned


def _write_frame(ws, df: pd.DataFrame, formats: dict[str, str]) -> None:
    """One frame onto one sheet: header, rows, formats, widths, frozen header.

    `index=False` is not an option here — it is written column by column, so
    there is no index column to omit.
    """
    columns = list(df.columns)
    for j, name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=j, value=str(name))
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    for i, (_, row) in enumerate(df.iterrows(), start=2):
        for j, name in enumerate(columns, start=1):
            value = row[name]
            if value is not None and not isinstance(value, (str, bool)) and pd.isna(value):
                value = None
            elif isinstance(value, (pd.Timestamp,)):
                value = value.to_pydatetime()
            elif hasattr(value, "item") and not isinstance(value, (str, bytes)):
                # numpy scalars: openpyxl writes the Python value, not a repr.
                value = value.item()
            cell = ws.cell(row=i, column=j, value=value)
            fmt = _NUMBER_FORMATS.get(formats.get(name, ""))
            if fmt and isinstance(value, (int, float)) and not isinstance(value, bool):
                cell.number_format = fmt

    # Auto-width from the rendered width of the header and the first rows —
    # every row would mean re-reading a sheet that can be long, and the rows
    # past the first screenful rarely set the width.
    for j, name in enumerate(columns, start=1):
        widest = len(str(name))
        for value in df[name].head(200):
            if value is None:
                continue
            widest = max(widest, len(f"{value:,.2f}") if isinstance(value, float) else len(str(value)))
        ws.column_dimensions[get_column_letter(j)].width = max(_MIN_WIDTH, min(_MAX_WIDTH, widest + 2))

    ws.freeze_panes = "A2"


def build_workbook(sheets: list[dict]) -> bytes:
    """Render `[{title, frame, formats}, ...]` in order as one .xlsx.

    Sheets arrive already ordered and already filtered — a sub-analysis that was
    skipped contributes no sheet here, and its reason is on Method instead.
    """
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for sheet in sheets:
        ws = wb.create_sheet(safe_sheet_name(sheet["title"], used))
        _write_frame(ws, sheet["frame"], sheet.get("formats") or {})
    if not wb.sheetnames:
        _write_frame(wb.create_sheet("Method"), pd.DataFrame({"Note": ["No sheets produced."]}), {})

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


# ── The sheets themselves (§17.5) ─────────────────────────────────────────────

def _summary_frame(results: list[dict]) -> pd.DataFrame:
    """Headline metrics: the compared metrics from the sub-analyses that ran,
    with their deltas kept numeric."""
    rows = []
    for result in results:
        df = result.get("frame")
        if df is None or df.empty:
            continue
        if result["key"] == "entity_metrics":
            for _, row in df.iterrows():
                rows.append({"Metric": row["Metric"], "Event": row.get("Event"),
                             "Baseline avg": row.get("Baseline avg"),
                             "Delta": row.get("Delta"), "% delta": row.get("% delta")})
        elif result["key"] == "event_vs_baseline" and "is total" in df.columns:
            totals = df[df["is total"] == True]  # noqa: E712
            if not len(totals):
                continue
            total = totals.iloc[0]
            for column in [c for c in df.columns if str(c).endswith("% delta")]:
                base = str(column)[: -len(" % delta")]
                rows.append({
                    "Metric": f"{base} (whole event window)",
                    "Event": total.get(f"{base} (event)"),
                    "Baseline avg": total.get(f"{base} (baseline avg)"),
                    "Delta": total.get(f"{base} delta"),
                    "% delta": total.get(column),
                })
    if not rows:
        return pd.DataFrame({"Metric": ["Nothing was compared against a baseline in this run."]})
    return pd.DataFrame(rows)


def _narrative_frame(narrative: str | None, findings: list[dict]) -> pd.DataFrame:
    """The written findings, one per row, so they survive as text (§17.5)."""
    rows = []
    for line in (narrative or "").splitlines():
        line = line.strip().lstrip("-*• ").strip()
        if line:
            rows.append({"Source": "Narrative", "Finding": line})
    for f in findings:
        for line in f["findings"]:
            rows.append({"Source": f["title"], "Finding": line})
    if not rows:
        rows = [{"Source": "Narrative", "Finding": "No findings were produced for this run."}]
    return pd.DataFrame(rows)


def _method_frame(table, roles, windows, results, extra_notes=None) -> pd.DataFrame:
    """What the system assumed, in the workbook rather than in the chat.

    Every sheet's provenance is here: the role each column filled and whether
    that was inferred or confirmed, both windows, the SQL behind each sheet, the
    rows it returned, and the reason for anything that did not run.
    """
    rows = [{"Section": "Source", "Item": "Table", "Detail": table, "Value": None}]

    for role, entry in roles.items():
        column = entry.get("column")
        rows.append({
            "Section": "Resolved role",
            "Item": role,
            "Detail": (f"{', '.join(column) if isinstance(column, list) else column} "
                       f"[{entry['source']}] — {entry['why']}").strip(),
            "Value": entry.get("confidence"),
        })

    event = windows.get("event") or {}
    rows.append({"Section": "Window", "Item": "Event window",
                 "Detail": f"{event.get('start')} to {event.get('end')}",
                 "Value": event.get("days")})
    for b in windows.get("baseline", []):
        rows.append({"Section": "Window", "Item": b["label"],
                     "Detail": f"{b['start']} to {b['end']}", "Value": None})
    rows.append({"Section": "Window", "Item": "How these were chosen",
                 "Detail": windows.get("note"), "Value": None})
    rows.append({"Section": "Window", "Item": "Data range",
                 "Detail": f"{windows.get('data_start')} to {windows.get('data_end')}",
                 "Value": None})

    for result in results:
        if result.get("skipped"):
            rows.append({"Section": "Sub-analysis", "Item": result["title"],
                         "Detail": f"SKIPPED — {result['skipped']}", "Value": None})
            continue
        rows.append({"Section": "Sub-analysis", "Item": result["title"],
                     "Detail": result["description"],
                     "Value": result.get("row_count")})
        rows.append({"Section": "Query", "Item": result["title"],
                     "Detail": result.get("sql"), "Value": None})
        if result.get("validation"):
            rows.append({"Section": "Query", "Item": f"{result['title']} — validation",
                         "Detail": result["validation"], "Value": None})

    for note in extra_notes or []:
        rows.append({"Section": "Note", "Item": "", "Detail": note, "Value": None})

    return pd.DataFrame(rows)


def build_event_workbook(table, roles, windows, results, findings, narrative,
                         extra_notes=None) -> bytes:
    """Assemble §17.5's workbook: Summary, Narrative, one sheet per
    sub-analysis that ran, then Method."""
    sheets = [
        {"title": "Summary", "frame": _summary_frame(results),
         "formats": {"Event": "amount", "Baseline avg": "amount",
                     "Delta": "amount", "% delta": "pct"}},
        {"title": "Narrative", "frame": _narrative_frame(narrative, findings)},
    ]
    for result in results:
        if result.get("skipped") or result.get("frame") is None:
            continue
        sheets.append({"title": result["title"], "frame": result["frame"],
                       "formats": result.get("formats") or {}})
    sheets.append({"title": "Method",
                   "frame": _method_frame(table, roles, windows, results, extra_notes)})
    return build_workbook(sheets)
